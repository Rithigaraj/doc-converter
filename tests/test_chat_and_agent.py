import io
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
import pymupdf
import openpyxl

from app.main import app
from agent.text_extractor import extract_document_text
from agent.llm_client import LLMProvider, set_llm_provider

client = TestClient(app)


class MockLLMProvider(LLMProvider):
    """Mock LLM Provider for deterministic unit testing without API credits."""

    def summarize(self, text: str, style: str = "detailed", length: str = "standard") -> str:
        return f"Mocked Summary ({style}, {length}): The document discusses {text[:50]}..."

    def chat_stream(self, messages: list[dict], document_context: str = ""):
        last_msg = messages[-1]["content"] if messages else ""
        yield f"Mock answer to '{last_msg}'. "
        yield f"Grounded context length is {len(document_context)}."


@pytest.fixture(autouse=True)
def setup_mock_llm():
    set_llm_provider(MockLLMProvider())


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    p = tmp_path / "test.docx"
    doc = Document()
    doc.add_heading("Contract Agreement", level=1)
    doc.add_paragraph("This agreement is entered into by Party A and Party B.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Duration"
    table.cell(1, 1).text = "12 Months"
    doc.save(str(p))
    return p


@pytest.fixture
def sample_pptx(tmp_path) -> Path:
    p = tmp_path / "test.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Q3 Financial Review"
    slide.placeholders[1].text = "Revenue grew by 24% year over year."
    prs.save(str(p))
    return p


@pytest.fixture
def sample_xlsx(tmp_path) -> Path:
    p = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Target", "Actual"])
    ws.append(["CAC", "100", "85"])
    wb.save(str(p))
    return p


def test_text_extraction_multi_format(sample_docx, sample_pptx, sample_xlsx, tmp_path):
    # 1. DOCX
    docx_doc = extract_document_text(sample_docx)
    assert "Contract Agreement" in docx_doc.text
    assert "Party A" in docx_doc.text
    assert "Duration" in docx_doc.text

    # 2. PPTX
    pptx_doc = extract_document_text(sample_pptx)
    assert "Q3 Financial Review" in pptx_doc.text
    assert "Revenue grew" in pptx_doc.text

    # 3. XLSX
    xlsx_doc = extract_document_text(sample_xlsx)
    assert "Metric" in xlsx_doc.text
    assert "CAC" in xlsx_doc.text

    # 4. HTML
    html_p = tmp_path / "test.html"
    html_p.write_text("<html><body><h1>Notice</h1><p>Confidential content.</p></body></html>", encoding="utf-8")
    html_doc = extract_document_text(html_p)
    assert "Notice" in html_doc.text
    assert "Confidential content" in html_doc.text


def test_chat_upload_endpoint(sample_docx):
    with open(sample_docx, "rb") as f:
        res = client.post(
            "/api/chat/upload",
            files={"file": ("contract.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    assert res.status_code == 200
    data = res.json()
    assert "doc_id" in data
    assert data["filename"] == "contract.docx"
    assert data["char_count"] > 0
    assert "Contract Agreement" in data["preview"]


def test_chat_summarize_endpoint(sample_docx):
    # Upload first
    with open(sample_docx, "rb") as f:
        up_res = client.post(
            "/api/chat/upload",
            files={"file": ("contract.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    doc_id = up_res.json()["doc_id"]

    # Summarize
    sum_res = client.post(
        "/api/chat/summarize",
        json={"doc_id": doc_id, "style": "bullet_points", "length": "brief"}
    )
    assert sum_res.status_code == 200
    data = sum_res.json()
    assert data["doc_id"] == doc_id
    assert "Mocked Summary" in data["summary"]
    assert "bullet_points" in data["summary"]


def test_chat_message_streaming(sample_docx):
    # Upload first
    with open(sample_docx, "rb") as f:
        up_res = client.post(
            "/api/chat/upload",
            files={"file": ("contract.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    doc_id = up_res.json()["doc_id"]

    # Send message
    res = client.post(
        "/api/chat/message",
        json={
            "session_id": "test_session",
            "doc_id": doc_id,
            "message": "What is the duration?",
            "history": []
        }
    )
    assert res.status_code == 200
    content = res.text
    assert "data: " in content
    assert "Mock answer" in content
    assert "[DONE]" in content


def test_export_summary_formats():
    for fmt in ["txt", "docx", "pdf"]:
        res = client.post(
            "/api/chat/export-summary",
            json={
                "title": "Contract Summary",
                "text": "1. Key Term: 12 months duration.\n2. Both parties agreed.",
                "format": fmt
            }
        )
        assert res.status_code == 200
        assert len(res.content) > 0
