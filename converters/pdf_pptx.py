import io
from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.util import Inches, Pt
import pymupdf

from app.utils.detect import is_scanned_pdf
from .base import BaseConverter, ConvertResult


class PdfToPptxConverter(BaseConverter):
    name = "pdf_to_pptx"

    def convert(self, input_path: Path, output_path: Path, options: dict[str, Any] | None = None) -> ConvertResult:
        options = options or {}
        warnings: list[str] = []
        try:
            if is_scanned_pdf(input_path):
                warnings.append(
                    "This appears to be a scanned/image-based PDF. Slide text extraction may be limited. "
                    "Pages rendered as high-resolution slide backgrounds."
                )

            doc = pymupdf.open(str(input_path))
            total_pages = len(doc)
            if total_pages == 0:
                doc.close()
                return ConvertResult(ok=False, error="PDF has no pages", warnings=warnings)

            prs = Presentation()
            # Standard widescreen 16:9 (13.33 x 7.5 inches) or 4:3 (10 x 7.5 inches)
            # Default to matching first page aspect ratio
            first_page = doc[0]
            rect = first_page.rect
            page_w, page_h = rect.width, rect.height
            if page_w > 0 and page_h > 0:
                aspect = page_w / page_h
                slide_height = Inches(7.5)
                slide_width = Inches(7.5 * aspect)
                prs.slide_width = slide_width
                prs.slide_height = slide_height

            blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

            for i, page in enumerate(doc):
                slide = prs.slides.add_slide(blank_layout)
                
                # Render page image at high DPI
                dpi = int(options.get("dpi", 150))
                pix = page.get_pixmap(dpi=dpi)
                img_bytes = pix.tobytes("png")
                
                slide.shapes.add_picture(
                    io.BytesIO(img_bytes),
                    left=0,
                    top=0,
                    width=prs.slide_width,
                    height=prs.slide_height
                )

                # Extract text and store in slide notes for searchability
                page_text = page.get_text("text").strip()
                if page_text:
                    notes_slide = slide.notes_slide
                    notes_text_frame = notes_slide.notes_text_frame
                    notes_text_frame.text = page_text

            doc.close()
            output_path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(output_path))
            warnings.append(f"Converted {total_pages} page(s) into PowerPoint slides with high-fidelity visuals and searchable slide notes.")
            return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "pymupdf+python-pptx", "pages": total_pages})
        except Exception as e:
            return ConvertResult(ok=False, error=f"PDF to PPTX conversion failed: {str(e)}", warnings=warnings)
