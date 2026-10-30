from pathlib import Path
from typing import Any
import openpyxl
import pdfplumber

from app.utils.detect import is_scanned_pdf
from .base import BaseConverter, ConvertResult


class PdfToXlsxConverter(BaseConverter):
    name = "pdf_to_xlsx"

    def convert(self, input_path: Path, output_path: Path, options: dict[str, Any] | None = None) -> ConvertResult:
        options = options or {}
        warnings: list[str] = []
        try:
            if is_scanned_pdf(input_path):
                warnings.append(
                    "This appears to be a scanned/image-based PDF. Table extraction may be incomplete. Consider OCR."
                )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Extracted Data"

            has_tables = False
            current_row = 1

            with pdfplumber.open(str(input_path)) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    tables = page.extract_tables() or []
                    if tables:
                        has_tables = True
                        for table in tables:
                            for row_data in table:
                                for col_idx, val in enumerate(row_data or []):
                                    ws.cell(row=current_row, column=col_idx + 1, value=str(val) if val is not None else "")
                                current_row += 1
                            current_row += 1  # Blank line between tables
                    else:
                        # Fallback: extract text lines and place in sheet
                        text = page.extract_text() or ""
                        if text.strip():
                            for line in text.splitlines():
                                parts = [p.strip() for p in line.split("\t") if p.strip()] or [line.strip()]
                                for col_idx, val in enumerate(parts):
                                    ws.cell(row=current_row, column=col_idx + 1, value=val)
                                current_row += 1
                            current_row += 1

            if not has_tables:
                warnings.append("No explicit grid tables detected; text lines extracted into spreadsheet.")
            else:
                warnings.append("Tables detected and extracted into spreadsheet.")

            output_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(output_path))
            warnings.append("Complex tables/charts/layouts require validation. Best-effort extraction.")
            return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "pdfplumber+openpyxl", "has_tables": has_tables})
        except Exception as e:
            return ConvertResult(ok=False, error=f"PDF to XLSX conversion failed: {str(e)}", warnings=warnings)
