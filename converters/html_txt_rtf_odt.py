import html
import os
import shutil
import subprocess
from pathlib import Path
from typing import Any
from bs4 import BeautifulSoup
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from app.config import settings, get_soffice_binary
from .base import BaseConverter, ConvertResult
from .gotenberg_client import GotenbergConverter


class HtmlTxtRtfOdtConverter(BaseConverter):
    name = "html_txt_rtf_odt"

    def convert(self, input_path: Path, output_path: Path, options: dict[str, Any] | None = None) -> ConvertResult:
        options = options or {}
        warnings: list[str] = []
        src = input_path.suffix.lower().lstrip(".")
        dst = output_path.suffix.lower().lstrip(".")
        output_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            # 1. TXT -> DOCX
            if src == "txt" and dst == "docx":
                doc = Document()
                content = input_path.read_text(encoding="utf-8", errors="ignore")
                for line in content.splitlines():
                    doc.add_paragraph(line)
                doc.save(str(output_path))
                warnings.append("Text converted to DOCX document.")
                return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "python-docx"})

            # 2. TXT -> PDF
            if src == "txt" and dst == "pdf":
                doc = SimpleDocTemplate(
                    str(output_path),
                    pagesize=letter,
                    rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
                )
                styles = getSampleStyleSheet()
                body_style = ParagraphStyle(
                    'TxtBody',
                    parent=styles['Normal'],
                    fontName='Helvetica',
                    fontSize=10,
                    leading=14
                )
                story = []
                content = input_path.read_text(encoding="utf-8", errors="ignore")
                for line in content.splitlines():
                    if line.strip():
                        safe_line = html.escape(line)
                        story.append(Paragraph(safe_line, body_style))
                    else:
                        story.append(Spacer(1, 8))
                if not story:
                    story.append(Paragraph("", body_style))
                doc.build(story)
                warnings.append("Text converted to PDF via ReportLab.")
                return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "reportlab"})

            # 3. HTML / HTM -> DOCX
            if src in {"html", "htm"} and dst == "docx":
                content = input_path.read_text(encoding="utf-8", errors="ignore")
                soup = BeautifulSoup(content, "html.parser")
                doc = Document()
                # Extract title
                title_tag = soup.find("title")
                if title_tag and title_tag.text:
                    doc.add_heading(title_tag.text.strip(), level=1)
                for element in soup.find_all(["h1", "h2", "h3", "p", "li"]):
                    text = element.get_text().strip()
                    if not text:
                        continue
                    if element.name == "h1":
                        doc.add_heading(text, level=1)
                    elif element.name == "h2":
                        doc.add_heading(text, level=2)
                    elif element.name == "h3":
                        doc.add_heading(text, level=3)
                    elif element.name == "li":
                        doc.add_paragraph(text, style='List Bullet')
                    else:
                        doc.add_paragraph(text)
                doc.save(str(output_path))
                warnings.append("HTML structure extracted to DOCX.")
                return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "bs4+python-docx"})

            # 4. HTML / HTM -> PDF
            if src in {"html", "htm"} and dst == "pdf":
                # Try Gotenberg HTML engine if available
                if GotenbergConverter.is_available():
                    ok, _ = GotenbergConverter.convert_to_pdf(input_path, output_path)
                    if ok:
                        warnings.append("HTML converted to PDF via Gotenberg.")
                        return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "gotenberg"})

                # Try LibreOffice
                soffice = get_soffice_binary()
                if soffice:
                    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_path.parent), str(input_path)]
                    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
                    produced = output_path.parent / f"{input_path.stem}.pdf"
                    if produced.exists():
                        if produced != output_path:
                            produced.replace(output_path)
                        warnings.append("HTML converted to PDF via LibreOffice.")
                        return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "libreoffice"})

                # Native fallback: ReportLab using parsed HTML
                content = input_path.read_text(encoding="utf-8", errors="ignore")
                soup = BeautifulSoup(content, "html.parser")
                doc = SimpleDocTemplate(str(output_path), pagesize=letter)
                styles = getSampleStyleSheet()
                story = []
                for element in soup.find_all(["h1", "h2", "h3", "p"]):
                    text = html.escape(element.get_text().strip())
                    if not text:
                        continue
                    if element.name == "h1":
                        story.append(Paragraph(f"<b>{text}</b>", styles['Heading1']))
                    elif element.name == "h2":
                        story.append(Paragraph(f"<b>{text}</b>", styles['Heading2']))
                    elif element.name == "h3":
                        story.append(Paragraph(f"<b>{text}</b>", styles['Heading3']))
                    else:
                        story.append(Paragraph(text, styles['Normal']))
                    story.append(Spacer(1, 6))
                if not story:
                    story.append(Paragraph(html.escape(soup.get_text()[:5000]), styles['Normal']))
                doc.build(story)
                warnings.append("HTML converted to PDF via layout parser.")
                return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "reportlab-html"})

            # 5. RTF / ODT -> PDF / DOCX
            if src in {"rtf", "odt"}:
                # Try Gotenberg
                if GotenbergConverter.is_available():
                    ok, _ = GotenbergConverter.convert_to_pdf(input_path, output_path)
                    if ok:
                        warnings.append(f"{src.upper()} converted via Gotenberg.")
                        return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "gotenberg"})

                # Try LibreOffice
                soffice = get_soffice_binary()
                if soffice:
                    target_fmt = "pdf" if dst == "pdf" else "docx"
                    cmd = [soffice, "--headless", "--convert-to", target_fmt, "--outdir", str(output_path.parent), str(input_path)]
                    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
                    produced = output_path.parent / f"{input_path.stem}.{target_fmt}"
                    if produced.exists():
                        if produced != output_path:
                            produced.replace(output_path)
                        warnings.append(f"{src.upper()} converted via LibreOffice.")
                        return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "libreoffice"})

                # Try pypandoc if available
                if shutil.which("pandoc") or settings.pandoc_path:
                    try:
                        import pypandoc
                        pypandoc.convert_file(str(input_path), to=dst, outputfile=str(output_path))
                        warnings.append("Converted via pandoc.")
                        return ConvertResult(ok=True, output_path=output_path, warnings=warnings, meta={"engine": "pandoc"})
                    except Exception:
                        pass

                return ConvertResult(
                    ok=False,
                    error=f"Conversion from {src.upper()} to {dst.upper()} requires Gotenberg (`docker compose up gotenberg`) or LibreOffice.",
                    warnings=warnings
                )

            return ConvertResult(ok=False, error=f"Unsupported format pair: {src} -> {dst}", warnings=warnings)
        except Exception as e:
            return ConvertResult(ok=False, error=f"Conversion error: {str(e)}", warnings=warnings)
